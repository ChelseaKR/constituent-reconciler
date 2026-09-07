"""Excel workbooks as a structured source, and the shapes that must be refused.

The reason this reader exists is that the manual step it removes -- export each
sheet to CSV before every run -- is where a column gets dropped or renamed. So
the suite's centre of gravity is not "can it read a cell". It is that a workbook
of the demo data resolves to exactly the record ids and review queue its CSV
does, and that every way a workbook can be *quietly* wrong is refused by name
instead of read as a blank.

Three of those quiet failures are real and specific to this format:

* A merged header cell reads as ``None`` in openpyxl's read-only mode --
  identical to an empty cell, and read-only worksheets do not expose merge
  ranges at all. Read as a blank, a merged header silently drops the column
  under it.
* A formula whose result Excel never cached also reads as ``None`` under
  ``data_only=True``. That is a value that was never computed being published
  as an empty one.
* A workbook's used range routinely extends past its last real row, so a
  reader that trusts it mints records with no fields -- people who do not
  exist.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from constituent_reconciler import excel, pipeline
from constituent_reconciler.cli import main
from constituent_reconciler.config import RecipeError, load_recipe
from constituent_reconciler.excel import WorkbookError
from constituent_reconciler.progress import ProgressEvent

EXAMPLES = Path(__file__).resolve().parents[1] / "examples" / "intake-demo"


def _sheet_from_csv(worksheet: object, csv_path: Path, *, blank_rows_above: int = 0) -> None:
    """Copy a CSV verbatim into a worksheet, values as text.

    Text, not typed values, so this helper isolates the reader from any date or
    number coercion: the parity test then measures the reader, not the fixture
    builder's choices.
    """

    for _ in range(blank_rows_above):
        worksheet.append([])  # type: ignore[attr-defined]
    with csv_path.open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.reader(handle):
            worksheet.append(row)  # type: ignore[attr-defined]


def _workbook_demo(tmp_path: Path, *, sheet_name: str = "Contacts", extra_sheets: int = 0) -> Path:
    """The bundled intake demo, with both CSV inputs re-saved as workbooks."""

    demo = tmp_path / "demo"
    demo.mkdir(parents=True)
    shutil.copy(EXAMPLES / "recipe.toml", demo / "recipe.toml")
    for name in ("existing", "incoming"):
        book = Workbook()
        first = book.active
        first.title = sheet_name
        _sheet_from_csv(first, EXAMPLES / f"{name}.csv")
        for index in range(extra_sheets):
            book.create_sheet(f"Notes{index + 1}")
        book.save(demo / f"{name}.xlsx")
    recipe = (demo / "recipe.toml").read_text(encoding="utf-8")
    recipe = recipe.replace('existing = "existing.csv"', 'existing = "existing.xlsx"')
    recipe = recipe.replace('incoming = "incoming.csv"', 'incoming = "incoming.xlsx"')
    recipe = recipe.replace("[mapping]", f'sheet = "{sheet_name}"\n\n[mapping]', 1)
    (demo / "recipe.toml").write_text(recipe, encoding="utf-8")
    return demo


def _csv_demo(tmp_path: Path) -> Path:
    demo = tmp_path / "csv-demo"
    demo.mkdir()
    for name in ("recipe.toml", "existing.csv", "incoming.csv"):
        shutil.copy(EXAMPLES / name, demo / name)
    return demo


def _run(recipe: Path, out_dir: Path) -> None:
    assert main(["run", "--config", str(recipe), "--out", str(out_dir)]) == 0


def _resolved_ids(out_dir: Path) -> list[str]:
    with (out_dir / "resolved.csv").open(newline="", encoding="utf-8") as handle:
        return [row[0] for row in csv.reader(handle)]


def _one_sheet(tmp_path: Path, rows: list[list[object]], *, name: str = "Contacts") -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = name
    for row in rows:
        sheet.append(row)
    path = tmp_path / "book.xlsx"
    book.save(path)
    return path


# -- the acceptance list from the issue ---------------------------------------


def test_workbook_demo_resolves_exactly_as_the_csv_demo(tmp_path: Path) -> None:
    """Same record ids and same review queue, byte for byte."""

    csv_out = tmp_path / "out-csv"
    xlsx_out = tmp_path / "out-xlsx"
    _run(_csv_demo(tmp_path) / "recipe.toml", csv_out)
    _run(_workbook_demo(tmp_path) / "recipe.toml", xlsx_out)

    assert _resolved_ids(xlsx_out) == _resolved_ids(csv_out)
    assert (xlsx_out / "review_queue.csv").read_bytes() == (
        csv_out / "review_queue.csv"
    ).read_bytes()


def test_a_named_sheet_that_does_not_exist_is_refused_by_name(tmp_path: Path) -> None:
    demo = _workbook_demo(tmp_path)
    recipe = (demo / "recipe.toml").read_text(encoding="utf-8")
    (demo / "recipe.toml").write_text(
        recipe.replace('sheet = "Contacts"', 'sheet = "Clients"'), encoding="utf-8"
    )
    with pytest.raises(RecipeError) as caught:
        _run(demo / "recipe.toml", tmp_path / "out")
    assert "'Clients'" in str(caught.value)
    assert "'Contacts'" in str(caught.value)
    assert not (tmp_path / "out" / "resolved.csv").exists()


def test_reordering_the_sheets_does_not_move_a_record_id(tmp_path: Path) -> None:
    """Sheet order is a property of the file; a named sheet is not."""

    first = _workbook_demo(tmp_path / "a", extra_sheets=2)
    _run(first / "recipe.toml", tmp_path / "out-a")

    reordered = _workbook_demo(tmp_path / "b", extra_sheets=2)
    for name in ("existing", "incoming"):
        path = reordered / f"{name}.xlsx"
        book = load_workbook(path)
        book.move_sheet("Contacts", offset=2)
        book.save(path)
    _run(reordered / "recipe.toml", tmp_path / "out-b")

    assert _resolved_ids(tmp_path / "out-b") == _resolved_ids(tmp_path / "out-a")


def test_a_merged_header_is_refused_not_read_as_blank(tmp_path: Path) -> None:
    """The case openpyxl's read-only mode cannot see for itself.

    The merge is placed at the *end* of the header row on purpose. A merged
    cell in the middle would also be caught by the blank-header refusal, so it
    would not prove the merge check does anything; a trailing one reads as a
    trailing blank, gets trimmed, and silently drops its column.
    """

    book = Workbook()
    sheet = book.active
    sheet.title = "Contacts"
    sheet.append(["First Name", "Last Name", "Email"])
    sheet.append(["Ada", "Lovelace", "ada@example.org"])
    sheet.merge_cells("B1:C1")
    path = tmp_path / "merged.xlsx"
    book.save(path)

    with pytest.raises(WorkbookError) as caught:
        excel.read_rows(path, sheet="Contacts")
    message = str(caught.value)
    assert "merged cell" in message
    assert "B1:C1" in message


def test_a_blank_header_inside_the_row_is_refused(tmp_path: Path) -> None:
    path = _one_sheet(tmp_path, [["First Name", None, "Email"], ["Ada", "Lovelace", "a@b.org"]])
    with pytest.raises(WorkbookError) as caught:
        excel.read_rows(path, sheet="Contacts")
    assert "blank header at B1" in str(caught.value)


def test_repeated_header_names_are_refused(tmp_path: Path) -> None:
    path = _one_sheet(tmp_path, [["Email", "Email"], ["a@b.org", "c@d.org"]])
    with pytest.raises(WorkbookError) as caught:
        excel.read_rows(path, sheet="Contacts")
    assert "repeats a column name: 'Email'" in str(caught.value)


def test_a_formula_with_no_saved_result_is_refused_not_read_as_empty(tmp_path: Path) -> None:
    """``data_only=True`` gives ``None`` for both an empty cell and this one."""

    path = _one_sheet(
        tmp_path,
        [
            ["First Name", "Last Name", "Email"],
            ["Ada", "Lovelace", "=CONCATENATE(A2,B2)"],
        ],
    )
    with pytest.raises(WorkbookError) as caught:
        excel.read_rows(path, sheet="Contacts")
    message = str(caught.value)
    assert "cell C2" in message
    assert "never saved" in message


def test_a_password_protected_workbook_says_so(tmp_path: Path) -> None:
    path = tmp_path / "locked.xlsx"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(WorkbookError) as caught:
        excel.read_rows(path)
    assert "password-protected" in str(caught.value)


def test_a_corrupt_workbook_is_refused_as_unreadable(tmp_path: Path) -> None:
    path = tmp_path / "junk.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(WorkbookError) as caught:
        excel.read_rows(path)
    assert "not a readable" in str(caught.value)


def test_a_header_row_past_the_end_is_refused_not_read_as_zero_rows(tmp_path: Path) -> None:
    path = _one_sheet(tmp_path, [["First Name"], ["Ada"]])
    with pytest.raises(WorkbookError) as caught:
        excel.read_rows(path, sheet="Contacts", header_row=9)
    assert "no row 9" in str(caught.value)


def test_an_empty_sheet_is_refused(tmp_path: Path) -> None:
    book = Workbook()
    book.active.title = "Contacts"
    path = tmp_path / "empty.xlsx"
    book.save(path)
    with pytest.raises(WorkbookError) as caught:
        excel.read_rows(path, sheet="Contacts")
    assert "is empty" in str(caught.value)


# -- value rendering ----------------------------------------------------------


def test_typed_cells_render_as_the_text_a_csv_would_have_held(tmp_path: Path) -> None:
    """Excel hands back types where a CSV hands back text.

    A postcode stored as a number must not become ``90210.0``; a date-formatted
    cell must render ISO-8601 so a consent date parses; a boolean must land on
    a recognized consent token rather than an unrecognized status, which would
    fail closed to withheld and quietly drop a person from the export.
    """

    path = _one_sheet(
        tmp_path,
        [
            ["Zip", "Ratio", "DOB", "Seen At", "Consent", "Missing"],
            [
                90210,
                1.5,
                dt.datetime(1975, 8, 14),
                dt.datetime(2026, 1, 15, 9, 30),
                True,
                None,
            ],
        ],
    )
    assert excel.read_rows(path, sheet="Contacts") == [
        {
            "Zip": "90210",
            "Ratio": "1.5",
            "DOB": "1975-08-14",
            "Seen At": "2026-01-15 09:30:00",
            "Consent": "true",
            "Missing": "",
        }
    ]


def test_trailing_blank_rows_are_dropped_and_interior_blanks_are_kept(tmp_path: Path) -> None:
    """A used range past the last real row is Excel's artifact, not data."""

    path = _one_sheet(
        tmp_path,
        [
            ["First Name", "Last Name"],
            ["Ada", "Lovelace"],
            [None, None],
            ["Grace", "Hopper"],
            [None, None],
            [None, None],
        ],
    )
    rows = excel.read_rows(path, sheet="Contacts")
    assert [row["First Name"] for row in rows] == ["Ada", "", "Grace"]


def test_header_row_offset_reads_the_row_the_recipe_names(tmp_path: Path) -> None:
    path = _one_sheet(
        tmp_path,
        [
            ["Client export 2026", None],
            ["First Name", "Last Name"],
            ["Ada", "Lovelace"],
        ],
    )
    assert excel.read_rows(path, sheet="Contacts", header_row=2) == [
        {"First Name": "Ada", "Last Name": "Lovelace"}
    ]


def test_an_unnamed_sheet_resolves_to_the_first_one(tmp_path: Path) -> None:
    book = Workbook()
    book.active.title = "Contacts"
    book.active.append(["First Name"])
    book.active.append(["Ada"])
    book.create_sheet("Later")
    path = tmp_path / "book.xlsx"
    book.save(path)
    assert excel.resolve_sheet_name(path, None) == "Contacts"


# -- routing, accounting and validate -----------------------------------------


def test_a_folder_walk_reads_a_workbook_beside_a_csv(tmp_path: Path) -> None:
    demo = tmp_path / "demo"
    demo.mkdir()
    shutil.copy(EXAMPLES / "recipe.toml", demo / "recipe.toml")
    shutil.copy(EXAMPLES / "existing.csv", demo / "existing.csv")

    folder = demo / "incoming"
    folder.mkdir()
    with (EXAMPLES / "incoming.csv").open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    header, body = rows[0], rows[1:]
    half = len(body) // 2
    with (folder / "batch-a.csv").open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows([header, *body[:half]])
    book = Workbook()
    book.active.title = "Contacts"
    for row in [header, *body[half:]]:
        book.active.append(row)
    book.save(folder / "batch-b.xlsx")

    recipe = (demo / "recipe.toml").read_text(encoding="utf-8")
    recipe = recipe.replace('incoming = "incoming.csv"', 'incoming = "incoming"')
    recipe = recipe.replace("[mapping]", 'sheet = "Contacts"\n\n[mapping]', 1)
    (demo / "recipe.toml").write_text(recipe, encoding="utf-8")

    out = tmp_path / "out"
    _run(demo / "recipe.toml", out)
    report = json.loads((out / "run_report.json").read_text(encoding="utf-8"))
    read = [Path(name).name for name in report["ingest"]["files_read"]]
    assert "batch-b.xlsx" in read
    assert not [skip for skip in report["ingest"]["files_skipped"] if "xlsx" in skip["path"]]


class _RecordingSink:
    """Keeps every progress event, so a stage's denominator can be asserted."""

    def __init__(self) -> None:
        self.events: list[ProgressEvent] = []

    def emit(self, event: ProgressEvent) -> None:
        self.events.append(event)


def test_a_workbook_is_not_counted_as_a_document_to_extract(tmp_path: Path) -> None:
    """Reading a workbook is parsing a structured file, not extracting a page.

    The routing test for this used to be ``kind != "csv"``. Left alone, it would
    have counted every workbook toward the extract stage's denominator and
    reported extraction progress against files the extractor never opens.
    """

    demo = _workbook_demo(tmp_path)
    sink = _RecordingSink()
    pipeline.run(load_recipe(demo / "recipe.toml"), progress=sink)
    assert all(event.stage != "extract" for event in sink.events)
    ingest = [event for event in sink.events if event.stage == "ingest"]
    assert ingest[0].total == 2


def test_validate_names_the_sheet_a_run_will_read(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    demo = _workbook_demo(tmp_path)
    assert main(["validate", "--config", str(demo / "recipe.toml")]) == 0
    out = capsys.readouterr().out
    assert "(sheet 'Contacts', header row 1)" in out


def test_validate_reports_a_missing_sheet_as_a_problem_not_a_traceback(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    demo = _workbook_demo(tmp_path)
    recipe = (demo / "recipe.toml").read_text(encoding="utf-8")
    (demo / "recipe.toml").write_text(
        recipe.replace('sheet = "Contacts"', 'sheet = "Nope"'), encoding="utf-8"
    )
    assert main(["validate", "--config", str(demo / "recipe.toml")]) != 0
    # Problems go to stderr, so a caller piping stdout still sees the refusal.
    assert "'Nope'" in capsys.readouterr().err


# -- recipe keys --------------------------------------------------------------


def _recipe_with(tmp_path: Path, input_extra: str) -> Path:
    path = tmp_path / "recipe.toml"
    path.write_text(
        "[input]\n"
        'incoming = "incoming.csv"\n'
        f"{input_extra}\n"
        "\n[mapping]\n"
        'first_name = "First Name"\n'
        'last_name = "Last Name"\n',
        encoding="utf-8",
    )
    return path


def test_header_row_must_be_a_whole_row_number(tmp_path: Path) -> None:
    with pytest.raises(RecipeError) as caught:
        load_recipe(_recipe_with(tmp_path, "header_row = 1.5"))
    assert "header_row" in str(caught.value)


def test_header_row_is_one_based(tmp_path: Path) -> None:
    with pytest.raises(RecipeError) as caught:
        load_recipe(_recipe_with(tmp_path, "header_row = 0"))
    assert "1-based" in str(caught.value)


def test_a_blank_sheet_name_is_refused(tmp_path: Path) -> None:
    with pytest.raises(RecipeError) as caught:
        load_recipe(_recipe_with(tmp_path, 'sheet = "   "'))
    assert "sheet" in str(caught.value)


def test_the_defaults_are_first_sheet_and_row_one(tmp_path: Path) -> None:
    recipe = load_recipe(_recipe_with(tmp_path, 'id_column = "id"'))
    assert recipe.sheet is None
    assert recipe.header_row == 1


def test_the_sheet_part_lookup_survives_a_package_without_rels(tmp_path: Path) -> None:
    """A workbook whose package omits the rels part falls back, it does not crash."""

    path = _one_sheet(tmp_path, [["First Name"], ["Ada"]])
    stripped = tmp_path / "stripped.xlsx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(stripped, "w") as target:
        for item in source.infolist():
            if item.filename == "xl/_rels/workbook.xml.rels":
                continue
            target.writestr(item, source.read(item.filename))
    assert excel._sheet_part(stripped, "Contacts") is None
